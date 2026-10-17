import json, re, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

S="/tmp/claude-0/-home-user/4d1cc429-040b-50d9-9b74-ae951c2cbd75/scratchpad/"
vids=json.load(open(S+"vids.json")); posts=json.load(open(S+"posts.json"))
ecap={}
for line in open("/home/user/hrcash/marketing/video/daftar-30-video-reels.md"):
    m=re.match(r"\| (E\d+) \| ([^|]+) \| ([^|]+) \| (.+) \|$",line.strip())
    if m: ecap[m.group(1)]=(m.group(2).strip(),m.group(4).strip())

K={"K01":("Tiap tanggal 25, HP bos ramai begini","chat-tanggal-25"),"K02":("Perlu install berapa aplikasi buat absen dan gajian?","nol-aplikasi"),"K03":("Tanggal 25 tidak lagi bikin pusing","kalender-gajian"),
"K04":("Titip absen itu sudah biasa?","titip-absen"),"K05":("Rekap absen sebulan? Jangan hitung manual","rekap-otomatis"),"K06":("Gaji dikomplain, kok kurang?","gaji-dikomplain"),
"K07":("Cuti dan lembur tanpa kertas","cuti-lembur"),"K08":("Kasbon lupa dicatat? Pusing nagihnya","kasbon"),"K09":("Buku absen, rekap manual, titip absen. Cukup sudah.","buku-absen"),"K10":("Hitung gaji berjam-jam?","hitung-gaji")}

inv=[]  # code, jenis, lokasi, hook, versi, durasi, caption, catatan
for code,(hook,slug) in K.items():
    inv.append((code,"Video 18 detik","marketing/video/animasi/hrcash-%s-%s-(wa|bersih).mp4"%(code,slug),hook,"WA / bersih","~18 dtk","%s Setuju? Tulis di komentar."%hook,"Versi WA: kontak 081234556757. Versi bersih: tanpa ajakan."))
for code,v in vids.items():
    if code.startswith("N"):
        n=int(code[1:]); folder="set-30" if n<=30 else "set-60"
        end="harga" if v["end"]=="harga" else "bersih"
        inv.append((code,"Video 18-19 detik","marketing/video/%s/%s/hrcash-%s-%s.mp4"%(folder,end,code,v["slug"]),v["hook"],"Harga + WA" if end=="harga" else "Bersih","~18-19 dtk","%s Setuju? Tulis di komentar."%v["hook"],"Rp349.000 + 081234556757 di penutup" if end=="harga" else "Tanpa harga/kontak"))
for code,v in vids.items():
    if code.startswith("E"):
        j,cap=ecap[code]
        inv.append((code,"Reels/TikTok 21-22 dtk (%s)"%j,"marketing/video/set-90/hrcash-%s-%s.mp4"%(code,v["slug"]),v["hook"],"Bersih","~21-22 dtk",cap,"Tanpa harga/kontak"))
hc=[("IG1","Carousel 3 gambar (IG 4:5)","marketing/instagram/ig-carousel-konten-1/slide-1..3.png","Katanya masuk jam 8. Buktinya mana?","","","Katanya sudah masuk jam 8, tapi buktinya tidak ada. Absen yang benar punya bukti: selfie, lokasi, jam otomatis. Pernah mengalami ini? Tulis di komentar. Mau tanya-tanya, WA 081234556757.","Konten pertama akun baru"),
("TT1","Foto 3 gambar (TikTok 9:16)","marketing/instagram/tiktok-konten-1/slide-1..3.png","Katanya masuk jam 8. Buktinya mana?","","","Katanya masuk jam 8, buktinya mana? Absen pakai bukti: selfie + lokasi + jam otomatis. Pernah kejadian? Komen ya.","Mode foto TikTok"),
("IG2","Gambar 1 (IG 4:5)","marketing/instagram/ig-post/konten-2-rekap-sampai-malam.png","Akhir bulan, rekap absen sampai malam","","","Jam berapa biasanya kamu selesai merekap absen akhir bulan? Jujur saja di komentar.",""),
("IG3","Gambar 1 (IG 4:5)","marketing/instagram/ig-post/konten-3-gaji-dikomplain.png","Gaji beda hitungan, karyawan protes","","","Selisih gaji sekecil apa pun bisa jadi debat. Pernah kejadian di tempatmu?",""),
("IG4","Gambar 1 (IG 4:5)","marketing/instagram/ig-post/konten-4-titip-absen.png","Satu orang absen, tiga nama tercatat","","","Titip absen itu rahasia umum. Pernah menemukan di tempat usahamu?",""),
("IG5","Gambar 1 (IG 4:5)","marketing/instagram/ig-post/konten-5-siapa-sudah-masuk.png","Jam 9 pagi. Siapa yang sudah masuk?","","","Sekarang jam berapa? Kamu tahu siapa saja yang sudah masuk hari ini?","")]
inv+=hc
for p in posts:
    folder="feed-30" if int(p["id"][1:3])<=30 else "feed-60"
    inv.append((p["id"][:3],"Post IG %d gambar"%p["n"],"marketing/instagram/%s/%s/slide-1..%d.png"%(folder,p["id"],p["n"]),p["hook"],"","",p["cap"],"Tema umum, tag: "+p["tag"].title()))

# ---------- jadwal ----------
E=[c for c in vids if c.startswith("E")]; E.sort()
edu=[c for c in E if int(c[1:])<=15]; pro=[c for c in E if 16<=int(c[1:])<=27]; pop=[c for c in E if int(c[1:])>=28]
seq=[]; 
pools=[edu[:],pro[:],pop[:]]
order=[0,0,1,0,1,0,1,2,0,1,0,1,0,2,1,0,1,0,1,2]  # 20 weekday videos
for o in order: seq.append(pools[o].pop(0))
# weekend tiktok-only: 8 video N bersih
Nb=[c for c in sorted(vids) if c.startswith("N") and vids[c]["end"]=="bersih"]
wk=[Nb[i] for i in (0,2,4,7,10,13,16,19)]
feed=["IG1","P01","IG2","P11","IG3","P04","IG4","P21","IG5","P31","P12","P41","P22","P51","P32","P42"]
feed=[f for f in feed]

wb=Workbook()
F="Arial"
def font(**k): return Font(name=F,size=k.pop("size",10),**k)
hdr_fill=PatternFill("solid",fgColor="17436A"); hdr_font=font(bold=True,color="FFFFFF")
thin=Side(style="thin",color="D0DAE6"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
wrap=Alignment(wrap_text=True,vertical="top")

# --- Panduan ---
ws=wb.active; ws.title="Panduan"
ws["A1"]="Konten Planner: HrCash dan Cash Impact (4 minggu)"; ws["A1"].font=font(bold=True,size=14)
rows=[("Tanggal mulai (isi sendiri)",datetime.date(2026,10,19),"Ubah tanggal ini, seluruh jadwal ikut bergeser. Disarankan hari Senin."),
("Jam unggah Reels / video IG","12:00","Kisaran jam ramai 11.00-13.00 dan 19.00-21.00. Ubah di kolom Jam pada sheet Jadwal."),
("Jam unggah IG feed (gambar)","19:00",""),
("Jam unggah TikTok","20:30","Video yang sama dengan Reels, caption sedikit disesuaikan.")]
ws["A3"]="Pengaturan"; ws["A3"].font=font(bold=True,size=11)
for i,(a,b,c) in enumerate(rows,start=4):
    ws.cell(i,1,a).font=font(); x=ws.cell(i,2,b); x.font=font(color="0000FF"); x.fill=PatternFill("solid",fgColor="FFF2B3"); ws.cell(i,3,c).font=font(color="555555")
ws["B4"].number_format="dd mmm yyyy"
ws["A9"]="Cara pakai"; ws["A9"].font=font(bold=True,size=11)
steps=["1. Buka sheet 'Jadwal 4 Minggu'. Tiap baris = satu unggahan (tanggal, jam, platform, kode konten).",
"2. Cari file lewat kolom 'Kode konten' di sheet 'Daftar Konten' (ada nama file/folder, hook, dan caption).",
"3. Setelah mengunggah, ubah kolom Status menjadi 'Terposting' (ada pilihan dropdown).",
"4. Ringkasan di bawah dan kolom 'Terpakai' di Daftar Konten menghitung otomatis.",
"5. Kolom berwarna kuning pada sheet Panduan boleh kamu ubah. Sel lain berisi rumus, jangan ditimpa.",
"6. Sheet 'Iklan Meta' berisi rencana uji iklan. Mulai setelah minggu ke-1 (akun sudah ada isi)."]
for i,t in enumerate(steps,start=10): ws.cell(i,1,t).font=font()
ws["A17"]="Peringatan penting"; ws["A17"].font=font(bold=True,size=11,color="C2410C")
warn=["Video lama V1-V7 (folder marketing/video) masih memuat harga Rp289.000. Jangan dipakai. Pakai hanya set K, N, E.",
"Video set 'harga' (N) menampilkan Rp349.000 dan nomor 081234556757. Pastikan harga masih berlaku saat diunggah.",
"Konten hukum (lembur, cuti, THR) bersifat informasi umum. THR (video E04) dan cuti (E03) belum diverifikasi sumbernya. Cek dulu sebelum dipakai.",
"Contoh angka di video dan gambar adalah ilustrasi. Nama karyawan di video adalah contoh."]
for i,t in enumerate(warn,start=18): ws.cell(i,1,t).font=font(); ws.cell(i,1).alignment=Alignment(wrap_text=True,vertical="top")
ws["A23"]="Ringkasan jadwal (otomatis)"; ws["A23"].font=font(bold=True,size=11)
summ=[("Total unggahan","=COUNTA('Jadwal 4 Minggu'!A5:A200)"),("IG Reels","=COUNTIF('Jadwal 4 Minggu'!E5:E200,\"IG Reels\")"),("IG Feed","=COUNTIF('Jadwal 4 Minggu'!E5:E200,\"IG Feed\")"),("TikTok video","=COUNTIF('Jadwal 4 Minggu'!E5:E200,\"TikTok\")"),("TikTok foto","=COUNTIF('Jadwal 4 Minggu'!E5:E200,\"TikTok Foto\")"),("Sudah terposting","=COUNTIF('Jadwal 4 Minggu'!L5:L200,\"Terposting\")"),("Belum terposting","=B24-B29")]
for i,(a,f) in enumerate(summ,start=24): ws.cell(i,1,a).font=font(); c=ws.cell(i,2,f); c.font=font(); c.alignment=Alignment(horizontal="left")
ws["A32"]="Hashtag dasar"; ws["A32"].font=font(bold=True,size=11)
ws["A33"]="HrCash: #absensikaryawan #hrumkm #payroll #umkmindonesia #pemilikusaha #tipsbisnis #hrd"; ws["A34"]="Umum: #tipskeuangan #produktivitas #dunia kerja #karier #selfimprovement #umkm"
for r in (33,34): ws.cell(r,1).font=font()
ws["A36"]="Sumber: jadwal dan frekuensi adalah usulan kerja (belum diuji ke akunmu). Jam unggah mengikuti kebiasaan umum dan perlu disesuaikan setelah melihat statistik akun."; ws["A36"].font=font(italic=True,color="555555")
ws.column_dimensions["A"].width=48; ws.column_dimensions["B"].width=16; ws.column_dimensions["C"].width=70

# --- Daftar Konten ---
wd=wb.create_sheet("Daftar Konten")
heads=["Kode konten","Jenis","File / folder","Hook / judul","Versi","Durasi","Caption (saran)","Catatan","Terpakai di jadwal"]
for j,h in enumerate(heads,1):
    c=wd.cell(1,j,h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=border
for i,r in enumerate(inv,start=2):
    for j,v in enumerate(r,1):
        c=wd.cell(i,j,v); c.font=font(); c.alignment=wrap; c.border=border
    c=wd.cell(i,9,f"=IF(COUNTIF('Jadwal 4 Minggu'!$G$5:$G$200,A{i})>0,\"Ya\",\"Belum\")"); c.font=font(); c.border=border
for col,w in zip("ABCDEFGHI",[12,26,52,44,12,12,60,34,12]): wd.column_dimensions[col].width=w
wd.freeze_panes="B2"; wd.auto_filter.ref=f"A1:I{len(inv)+1}"

# --- Jadwal ---
wj=wb.create_sheet("Jadwal 4 Minggu",1)
wj["A1"]="Jadwal unggah 4 minggu"; wj["A1"].font=font(bold=True,size=14)
wj["A2"]="Tanggal mengikuti 'Tanggal mulai' di sheet Panduan. Hook dan caption terisi otomatis dari sheet Daftar Konten."; wj["A2"].font=font(italic=True,color="555555")
hj=["No","Tanggal","Hari","Jam","Platform","Format","Kode konten","Hook / judul","Tujuan","Caption (saran)","Catatan","Status"]
for j,h in enumerate(hj,1):
    c=wj.cell(4,j,h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=border
rows=[]  # (offset,jam_ref,platform,format,kode,tujuan,catatan)
vi=0; wi=0; fi=0
for day in range(28):
    wd_idx=day%7
    if wd_idx<5:
        code=seq[vi]; vi+=1
        j="E"+code[1:]; tuj="Edukasi" if int(code[1:])<=15 else ("Promosi" if int(code[1:])<=27 else "Interaksi")
        rows.append((day,"=Panduan!$B$5","IG Reels","Video 9:16",code,tuj,"Unggah ke Reels"))
        rows.append((day,"=Panduan!$B$7","TikTok","Video 9:16",code,tuj,"Video sama, caption disesuaikan"))
    else:
        code=wk[wi]; wi+=1
        rows.append((day,"=Panduan!$B$7","TikTok","Video 9:16",code,"Promosi ringan","Khusus TikTok akhir pekan"))
    if wd_idx in (0,2,4,6):
        code=feed[fi]; fi+=1
        tuj="Edukasi/Interaksi" if code.startswith("P") else "Masalah owner"
        rows.append((day,"=Panduan!$B$6","IG Feed","Gambar 4:5" if code!="IG1" else "Carousel 3 gambar",code,tuj,"Unggah di feed"))
    if day==2:
        rows.append((day,"=Panduan!$B$7","TikTok Foto","Foto 3 slide 9:16","TT1","Masalah owner","Mode foto TikTok"))
# sort by day then time order Reels(12) feed(19) tiktok(20:30)
order_p={"IG Reels":0,"IG Feed":1,"TikTok":2,"TikTok Foto":2}
rows.sort(key=lambda r:(r[0],order_p[r[2]]))
for i,r in enumerate(rows,start=5):
    n=i-4
    wj.cell(i,1,n); wj.cell(i,2,f"=Panduan!$B$4+{r[0]}")
    wj.cell(i,3,f'=CHOOSE(WEEKDAY(B{i},2),"Senin","Selasa","Rabu","Kamis","Jumat","Sabtu","Minggu")')
    wj.cell(i,4,r[1]); wj.cell(i,5,r[2]); wj.cell(i,6,r[3]); wj.cell(i,7,r[4])
    wj.cell(i,8,f"=IFERROR(INDEX('Daftar Konten'!$D:$D,MATCH(G{i},'Daftar Konten'!$A:$A,0)),\"\")")
    wj.cell(i,9,r[5])
    wj.cell(i,10,f"=IFERROR(INDEX('Daftar Konten'!$G:$G,MATCH(G{i},'Daftar Konten'!$A:$A,0)),\"\")")
    wj.cell(i,11,r[6]); wj.cell(i,12,"Belum")
    for j in range(1,13):
        c=wj.cell(i,j); c.font=font(); c.alignment=wrap; c.border=border
    wj.cell(i,2).number_format="dd mmm yyyy"; wj.cell(i,4).number_format="hh:mm"
    wj.cell(i,12).font=font(color="0000FF")
last=4+len(rows)
dv=DataValidation(type="list",formula1='"Belum,Terjadwal,Terposting"',allow_blank=False); wj.add_data_validation(dv); dv.add(f"L5:L{last+60}")
for col,w in zip("ABCDEFGHIJKL",[6,13,10,8,13,18,12,46,18,60,30,13]): wj.column_dimensions[col].width=w
wj.freeze_panes="C5"; wj.auto_filter.ref=f"A4:L{last}"
# warna per platform
fills={"IG Reels":"FDE9F1","IG Feed":"FFF4D6","TikTok":"E6F2FB","TikTok Foto":"E6F2FB"}
for i,r in enumerate(rows,start=5):
    for j in (5,):
        wj.cell(i,j).fill=PatternFill("solid",fgColor=fills[r[2]])

# --- Iklan Meta ---
wm=wb.create_sheet("Iklan Meta")
wm["A1"]="Rencana uji iklan Meta Ads (usulan)"; wm["A1"].font=font(bold=True,size=14)
wm["A2"]="Mulai setelah minggu ke-1. Angka budget diisi sendiri (sel kuning). Tidak ada angka hasil yang dijanjikan."; wm["A2"].font=font(italic=True,color="555555")
wm["A4"]="Budget per hari (isi sendiri)"; wm["A4"].font=font(bold=True); wm["B4"]=None; wm["B4"].fill=PatternFill("solid",fgColor="FFF2B3"); wm["B4"].font=font(color="0000FF"); wm["B4"].number_format='"Rp"#,##0'
wm["C4"]="Lama uji per tahap (hari)"; wm["C4"].font=font(bold=True); wm["D4"]=5; wm["D4"].fill=PatternFill("solid",fgColor="FFF2B3"); wm["D4"].font=font(color="0000FF")
wm["A5"]="Total budget per tahap"; wm["A5"].font=font(bold=True); wm["B5"]='=IF(B4="","",B4*D4)'; wm["B5"].number_format='"Rp"#,##0'; wm["B5"].font=font()
hm=["Tahap","Minggu","Tujuan","Kreatif (kode konten)","Alasan","Tujuan kampanye di Meta","Audiens (usulan)","Yang dipantau","Keputusan"]
for j,h in enumerate(hm,1):
    c=wm.cell(7,j,h); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=border
ads=[("1","Minggu 1","Organik dulu","Semua jadwal organik","Mengisi akun dan melihat konten yang disukai audiens","-","-","Jangkauan, simpan, bagikan, komentar","Catat 3 konten terbaik"),
("2","Minggu 2","Uji hook","N33, N49, N51","Tiga hook berbeda dengan harga Rp349.000 dan kontak WA","Pesan (WhatsApp)","Pemilik usaha kecil, usia 25-55, Indonesia, minat bisnis/UMKM","Tayangan 3 detik, CTR, biaya per pesan","Matikan yang terendah, lanjutkan 2"),
("3","Minggu 3","Uji pesan","N13, N46, N54","Fokus kecepatan, risiko data, tanpa HRD","Pesan (WhatsApp)","Sama seperti tahap 2","Tayangan 3 detik, CTR, biaya per pesan","Bandingkan dengan tahap 2"),
("4","Minggu 4","Skala pemenang","2 kreatif terbaik tahap 2-3","Naikkan budget bertahap pada yang konsisten","Pesan (WhatsApp)","Perluas minat dan wilayah","Biaya per pesan, jumlah pesan, kualitas prospek","Naikkan 20-30% bila stabil"),
("Cadangan","Kapan saja","Retargeting","E16, E17, E19 (versi bersih)","Menghangatkan yang sudah menonton","Interaksi / Pesan","Penonton video 50%+","Frekuensi, CTR","Hentikan jika frekuensi tinggi")]
for i,r in enumerate(ads,start=8):
    for j,v in enumerate(r,1):
        c=wm.cell(i,j,v); c.font=font(); c.alignment=wrap; c.border=border
for col,w in zip("ABCDEFGHI",[12,12,18,30,40,22,40,34,30]): wm.column_dimensions[col].width=w
wm["A15"]="Catatan: kreatif harga (set N 'harga') memuat Rp349.000 dan 081234556757. Selalu cek kebijakan iklan Meta sebelum menayangkan, terutama untuk klaim hukum (lembur, THR)."; wm["A15"].font=font(italic=True,color="555555")

out="/home/user/hrcash/marketing/Konten-Planner-4-Minggu.xlsx"
wb.save(out); print(out,len(rows),len(inv))
